from __future__ import annotations

import argparse
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

DEFAULT_INPUT = Path(__file__).resolve().with_name("EPI Database_CHDN.xlsx")
DEFAULT_OUTPUT = Path(__file__).resolve().with_name("CHDN_EPI_clean.xlsx")
DEFAULT_SOURCE_SHEET = "EPI-Child"
DEFAULT_TARGET_SHEET = "EPI-Child"
DEFAULT_PREGNANCY_SOURCE_SHEET = "EPI-Pregnancy"
DEFAULT_PREGNANCY_TARGET_SHEET = "EPI-Pregnancy"

DOSE_AUTOFILL_MAPPINGS = [
    ("BCG reporting month", "BCG other (Y/N)", "BCG Age ", "BCG Source"),
    ("OPV first time dose reporting month", "OPV first time dose other", "OPV1 Age ", "OPV1 Source"),
    ("OPV second time dose reporting month", "OPV second time dose other", "OPV2 Age", "OPV2 Source"),
    ("OPV third time dose reporting month", "OPV third time dose other", "OPV3 Age", "OPV3 Source"),
    ("Penta first time dose reporting month", "Penta first time dose other", "Penta1 Age", "Penta1 Source"),
    ("Penta second time dose reporting month", "Penta second time dose other", "Penta2 Age", "Penta2 Source"),
    ("Penta third time dose reporting month", "Penta third time dose other", "Penta3 Age", "Penta3 Source"),
    ("MMR first time dose reporting month", "MMR first time dose other", "MMR1 Age", "MMR1 Source"),
    ("MMR second time dose reporting month", "MMR second time dose other", "MMR2 Age", "MMR2 source"),
    ("JE first time dose reporting month", "JE first time dose other", "JE1 Age", "JE1 Source"),
    ("JE second time dose reporting month", "JE second time dose other", "JE2 Age", "JE2 Source"),
    ("IPV_reporting_month", "IPV dose other", "IPV Age", "IPV Source"),
]

COMPLETE_COLUMNS = [
    "CompleteInQ4 2024",
    "CompleteInQ1 2025",
    "CompleteInQ2 2025",
    "CompleteInQ3 2025",
    "CompleteInQ4 2025",
    "CompleteInQ1 2026",
    "CompleteInQ2 2026",
    "CompleteInQ3 2026",
    "CompleteInQ4 2026",
]

QUARTER_WINDOWS = {
    "CompleteInQ4 2024": ("Q4", date(2024, 9, 21), date(2024, 12, 20)),
    "CompleteInQ1 2025": ("Q1", date(2024, 12, 21), date(2025, 3, 20)),
    "CompleteInQ2 2025": ("Q2", date(2025, 3, 21), date(2025, 6, 20)),
    "CompleteInQ3 2025": ("Q3", date(2025, 6, 21), date(2025, 9, 20)),
    "CompleteInQ4 2025": ("Q4", date(2025, 9, 21), date(2025, 12, 20)),
    "CompleteInQ1 2026": ("Q1", date(2025, 12, 21), date(2026, 3, 20)),
    "CompleteInQ2 2026": ("Q2", date(2026, 3, 21), date(2026, 6, 20)),
    "CompleteInQ3 2026": ("Q3", date(2026, 6, 21), date(2026, 9, 20)),
    "CompleteInQ4 2026": ("Q4", date(2026, 9, 21), date(2026, 12, 20)),
}


def _normalize_col_name(name: str) -> str:
    return "".join(ch for ch in str(name).strip().lower() if ch.isalnum())


def _find_dob_col_idx(ws) -> int | None:
    aliases = {"dob", "dateofbirth", "birthdate", "datebirth"}
    for col_idx in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col_idx).value
        if header is None:
            continue
        norm = _normalize_col_name(header)
        if ("dob" in norm) or (norm in aliases) or ("dateofbirth" in norm):
            return col_idx
    if ws.max_column >= 11:
        return 11
    return None


def _header_col_map(ws) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col_idx).value
        if header is None:
            continue
        mapping[_normalize_col_name(header)] = col_idx
    return mapping


def _get_or_add_column(ws, header_name: str, header_map: dict[str, int]) -> int:
    key = _normalize_col_name(header_name)
    existing = header_map.get(key)
    if existing is not None:
        return existing
    new_idx = ws.max_column + 1
    ws.cell(row=1, column=new_idx, value=header_name)
    header_map[key] = new_idx
    return new_idx


def _is_blank(value) -> bool:
    return value is None or str(value).strip() == ""


def _to_date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
        parsed = pd.to_datetime(value, errors="coerce")
        if pd.isna(parsed):
            return None
        return parsed.date()
    if isinstance(value, (int, float)):
        try:
            parsed = pd.to_datetime(value, unit="D", origin="1899-12-30", errors="coerce")
            if pd.isna(parsed):
                return None
            return parsed.date()
        except Exception:
            return None
    return None


def _is_yes(value) -> bool:
    return str(value).strip().upper() == "YES"


def _is_no(value) -> bool:
    return str(value).strip().upper() == "NO"


def _datedif_months(start_date: date, end_date: date) -> int:
    months = (end_date.year - start_date.year) * 12 + (end_date.month - start_date.month)
    if end_date.day < start_date.day:
        months -= 1
    return months


def _apply_age_source_autofill(input_file: Path, source_sheet_name: str) -> tuple[Path, dict]:
    wb = load_workbook(input_file)
    if source_sheet_name not in wb.sheetnames:
        raise ValueError(f"Source sheet not found for autofill: {source_sheet_name}")

    ws = wb[source_sheet_name]
    dob_col_idx = _find_dob_col_idx(ws)
    if dob_col_idx is None:
        raise ValueError("DOB column not found in EPI-Child sheet.")

    header_map = _header_col_map(ws)
    stats = {"age_filled": 0, "source_filled": 0, "age_columns_added": 0, "source_columns_added": 0}

    prepared_mappings: list[tuple[int, int, int, int]] = []
    for dose_name, other_name, age_name, source_name in DOSE_AUTOFILL_MAPPINGS:
        dose_idx = header_map.get(_normalize_col_name(dose_name))
        other_idx = header_map.get(_normalize_col_name(other_name))
        if dose_idx is None or other_idx is None:
            continue
        age_key = _normalize_col_name(age_name)
        source_key = _normalize_col_name(source_name)
        age_was_missing = age_key not in header_map
        source_was_missing = source_key not in header_map
        age_idx = _get_or_add_column(ws, age_name, header_map)
        source_idx = _get_or_add_column(ws, source_name, header_map)
        if age_was_missing:
            stats["age_columns_added"] += 1
        if source_was_missing:
            stats["source_columns_added"] += 1
        prepared_mappings.append((dose_idx, other_idx, age_idx, source_idx))

    if not prepared_mappings:
        raise ValueError("No dose/other column pairs found for autofill mapping.")

    max_col = ws.max_column
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=max_col, values_only=False):
        row_values = [cell.value for cell in row]
        if not any(not _is_blank(value) for value in row_values):
            continue
        dob_cell = row[dob_col_idx - 1]
        dob_date = _to_date(dob_cell.value)
        for dose_idx, other_idx, age_idx, source_idx in prepared_mappings:
            dose_cell = row[dose_idx - 1]
            other_cell = row[other_idx - 1]
            age_cell = row[age_idx - 1] if age_idx <= len(row) else None
            source_cell = row[source_idx - 1] if source_idx <= len(row) else None
            dose_raw = dose_cell.value
            other_raw = other_cell.value
            if age_cell is not None and _is_blank(age_cell.value):
                try:
                    if _is_yes(other_raw):
                        age_cell.value = 999
                        stats["age_filled"] += 1
                    else:
                        dose_date = _to_date(dose_raw)
                        if dob_date is not None and dose_date is not None:
                            age_cell.value = _datedif_months(dob_date, dose_date)
                            stats["age_filled"] += 1
                        else:
                            age_cell.value = 1111
                            stats["age_filled"] += 1
                except Exception:
                    age_cell.value = ""
            if source_cell is not None and _is_blank(source_cell.value):
                try:
                    if (not _is_blank(dose_raw)) and _is_no(other_raw):
                        source_cell.value = "CHDN"
                        stats["source_filled"] += 1
                    elif _is_yes(other_raw):
                        source_cell.value = "Other"
                        stats["source_filled"] += 1
                    elif _is_blank(dose_raw) and _is_no(other_raw):
                        source_cell.value = "Not Received Yet"
                        stats["source_filled"] += 1
                    else:
                        source_cell.value = ""
                except Exception:
                    source_cell.value = ""

    temp_dir = Path(tempfile.mkdtemp(prefix="chdn_clean_copy_"))
    temp_input = temp_dir / input_file.name
    wb.save(temp_input)
    return temp_input, stats


def _find_children_code_column(columns) -> int | None:
    aliases = {"childrencodetccode", "childrencode", "tccode", "tcode"}
    for idx, col in enumerate(columns, start=1):
        if _normalize_col_name(col) in aliases:
            return idx
    return None


def _find_pregnance_code_column(columns) -> int | None:
    aliases = {"pregnancecode", "pregnancycode", "pregnance_code", "pregnancy_code", "pwcode", "pw_code"}
    norm_aliases = {a.replace("_", "") for a in aliases}
    for idx, col in enumerate(columns, start=1):
        if _normalize_col_name(col) in norm_aliases:
            return idx
    return None


def _find_header_index(header_index_map: dict[str, int], aliases: list[str]) -> int | None:
    for alias in aliases:
        if alias in header_index_map:
            return header_index_map[alias]
    normalized_aliases = {_normalize_col_name(alias) for alias in aliases}
    for header_name, index in header_index_map.items():
        if _normalize_col_name(header_name) in normalized_aliases:
            return index
    return None


def _row_value(row_values, header_index_map: dict[str, int], aliases: list[str]):
    index = _find_header_index(header_index_map, aliases)
    if index is None:
        return ""
    return row_values[index]


def _presence_ok(row_values, header_index_map: dict[str, int], dose_aliases: list[str], other_aliases: list[str]) -> bool:
    dose_value = _row_value(row_values, header_index_map, dose_aliases)
    other_value = _row_value(row_values, header_index_map, other_aliases)
    if not _is_blank(dose_value):
        return True
    if not _is_blank(other_value):
        return True
    return False


def _date_source_ok(row_values, header_index_map: dict[str, int], date_aliases: list[str], source_aliases: list[str], start: date, end: date) -> bool:
    row_date = _to_date(_row_value(row_values, header_index_map, date_aliases))
    if row_date is None or not (start <= row_date <= end):
        return False
    source_value = str(_row_value(row_values, header_index_map, source_aliases)).strip().upper()
    if source_value in {"CHDN", "CHDN/OTHER", "OTHER", ""}:
        return True
    return False


def _completion_value(row_values, column_name: str, header_index_map: dict[str, int]) -> str:
    quarter_label, start_date, end_date = QUARTER_WINDOWS[column_name]
    year_label = column_name.rsplit(" ", 1)[-1]
    u1_presence_rules = [
        (["BCG dose", "BCG"], ["BCG other (Y/N)", "BCG other", "BCG other yn"]),
        (["OPV first time dose", "OPV first timeadose", "OPV1 dose", "OPV1"], ["OPV first time dose other", "OPV1 other", "OPV1 other yn"]),
        (["OPV second time dose", "OPV2 dose", "OPV2"], ["OPV second time dose other", "OPV2 other", "OPV2 other yn"]),
        (["OPV third time dose", "OPV3 dose", "OPV3"], ["OPV third time dose other", "OPV3 other", "OPV3 other yn"]),
        (["Penta first time dose", "Penta1 dose", "Penta1"], ["Penta first time dose other", "Penta1 other", "Penta1 other yn"]),
        (["Penta second time dose", "Penta2 dose", "Penta2"], ["Penta second time dose other", "Penta2 other", "Penta2 other yn"]),
        (["Penta third time dose", "Penta3 dose", "Penta3"], ["Penta third time dose other", "Penta3 other", "Penta3 other yn"]),
        (["MMR first time dose", "MMR first timeadose", "MMR1 dose", "MMR1"], ["MMR first time dose other", "MMR1 other", "MMR1 other yn"]),
    ]
    u5_presence_rules = u1_presence_rules[1:]
    u1_date_source_rules = [
        (["BCG reporting month", "BCG report month"], ["BCG Source", "BCG source", "BCG src"]),
        (["OPV first time dose reporting month", "OPV1 reporting month"], ["OPV1 Source", "OPV1 source", "OPV1 src"]),
        (["OPV second time dose reporting month", "OPV2 reporting month"], ["OPV2 Source", "OPV2 source", "OPV2 src"]),
        (["OPV third time dose reporting month", "OPV3 reporting month"], ["OPV3 Source", "OPV3 source", "OPV3 src"]),
        (["Penta first time dose reporting month", "Penta1 reporting month"], ["Penta1 Source", "Penta1 source", "Penta1 src"]),
        (["Penta second time dose reporting month", "Penta2 reporting month"], ["Penta2 Source", "Penta2 source", "Penta2 src"]),
        (["Penta third time dose reporting month", "Penta3 reporting month"], ["Penta3 Source", "Penta3 source", "Penta3 src"]),
        (["MMR first time dose reporting month", "MMR1 reporting month"], ["MMR1 Source", "MMR1 source", "MMR1 src"]),
    ]
    u5_date_source_rules = u1_date_source_rules[1:]
    age_value = _row_value(row_values, header_index_map, ["Age at first visit", "Age", "Age in months", "Age at first visit "])
    try:
        age_num = float(age_value)
    except (TypeError, ValueError):
        age_num = None
    if age_num is not None and age_num <= 11:
        has_any_presence = any(_presence_ok(row_values, header_index_map, dose_aliases, other_aliases) for dose_aliases, other_aliases in u1_presence_rules)
        has_any_date_source = any(_date_source_ok(row_values, header_index_map, date_aliases, source_aliases, start_date, end_date) for date_aliases, source_aliases in u1_date_source_rules)
        if has_any_presence and has_any_date_source:
            return f"U1 complete in {quarter_label}_{year_label}"
    if age_num is not None and 11 < age_num <= 59:
        has_any_presence = any(_presence_ok(row_values, header_index_map, dose_aliases, other_aliases) for dose_aliases, other_aliases in u5_presence_rules)
        has_any_date_source = any(_date_source_ok(row_values, header_index_map, date_aliases, source_aliases, start_date, end_date) for date_aliases, source_aliases in u5_date_source_rules)
        if has_any_presence and has_any_date_source:
            return f"1-5 complete in {quarter_label}_{year_label}"
    return ""


def _warn_code_quality(sheet_or_values, code_col_idx_or_label, code_label: str | None = None, row_numbers: list[int] | None = None) -> None:
    if code_label is None and row_numbers is None:
        code_values = sheet_or_values
        label = code_col_idx_or_label
    else:
        sheet = sheet_or_values
        code_col_idx = code_col_idx_or_label
        code_values = [sheet.cell(row=row_num, column=code_col_idx).value for row_num in row_numbers or []]
        label = code_label
    normalized_codes = [str(value).strip() for value in code_values if not _is_blank(value)]
    missing_count = sum(1 for value in code_values if _is_blank(value))
    duplicate_codes = sorted({value for value in normalized_codes if normalized_codes.count(value) > 1})
    if missing_count > 0:
        print(f"WARNING: {label} missing. Missing count = {missing_count}")
    else:
        print(f"No missing {label}")
    if duplicate_codes:
        print(f"WARNING: duplication of {label} ({','.join(duplicate_codes)})")


def _get_kept_rows(values_sheet, data_start_row: int, max_col: int) -> tuple[list[int], int]:
    kept_rows: list[int] = []
    removed_count = 0
    max_row = values_sheet.max_row
    for row_num, row in enumerate(values_sheet.iter_rows(min_row=data_start_row, max_row=max_row, values_only=False), start=data_start_row):
        row_values = [cell.value for cell in row[:max_col]]
        if any(not _is_blank(value) for value in row_values):
            kept_rows.append(row_num)
        else:
            removed_count += 1
    return kept_rows, removed_count


def build_clean_sheet(
    input_file: Path | str,
    output_file: Path | str,
    source_sheet_name: str = "EPI-Child",
    target_sheet_name: str = "EPI-Child",
    pregnancy_source_sheet_name: str = "EPI-Pregnancy",
    pregnancy_target_sheet_name: str = "EPI-Pregnancy",
) -> int:
    input_path = Path(input_file).resolve()
    output_path = Path(output_file).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if not input_path.exists():
        raise FileNotFoundError(f"Input workbook not found: {input_path}")

    try:
        prefilled_input, stats = _apply_age_source_autofill(input_path, source_sheet_name)
    except Exception as exc:
        print(f"ERROR: Failed to prefill age/source values: {exc}")
        return 2

    print("Autofill summary (EPI-Child):")
    print(f"  Age values filled: {stats['age_filled']}")
    print(f"  Source values filled: {stats['source_filled']}")
    print(f"  Age columns added: {stats['age_columns_added']}")
    print(f"  Source columns added: {stats['source_columns_added']}")

    wb = load_workbook(prefilled_input, data_only=True, read_only=False)
    if source_sheet_name not in wb.sheetnames:
        raise ValueError(f"Source sheet not found: {source_sheet_name}")

    child_ws = wb[source_sheet_name]
    max_col = child_ws.max_column
    source_headers = [child_ws.cell(row=1, column=col_idx).value for col_idx in range(1, max_col + 1)]
    code_col_idx = _find_children_code_column(source_headers)
    if code_col_idx is None:
        raise ValueError("Could not find children_code(T/C Code) column in the source sheet")

    child_rows_to_keep: list[list[object]] = []
    child_code_values: list[object] = []
    child_removed_count = 0

    for row_num in range(2, child_ws.max_row + 1):
        row_values = [child_ws.cell(row=row_num, column=col_idx).value for col_idx in range(1, max_col + 1)]
        if any(not _is_blank(value) for value in row_values):
            child_rows_to_keep.append(row_values)
            child_code_values.append(row_values[code_col_idx - 1])
        else:
            child_removed_count += 1

    if child_removed_count > 0:
        print(f"INFO: Removed {child_removed_count} formula-only rows from {source_sheet_name}")

    _warn_code_quality(child_code_values, "children_code")

    out_wb = Workbook()
    child_out_ws = out_wb.active
    child_out_ws.title = target_sheet_name
    child_out_ws.append(list(source_headers) + list(COMPLETE_COLUMNS))

    header_index_map = {header_name: idx for idx, header_name in enumerate(source_headers)}
    for row_values in child_rows_to_keep:
        completed_values = {col_name: _completion_value(row_values, col_name, header_index_map) for col_name in COMPLETE_COLUMNS}
        child_out_ws.append(row_values + [completed_values[col_name] for col_name in COMPLETE_COLUMNS])

    if pregnancy_source_sheet_name in wb.sheetnames:
        preg_ws = wb[pregnancy_source_sheet_name]
        preg_max_col = preg_ws.max_column
        preg_header_row = 2
        preg_headers = [preg_ws.cell(row=preg_header_row, column=col_idx).value for col_idx in range(1, preg_max_col + 1)]

        preg_rows_to_keep, preg_removed_count = _get_kept_rows(preg_ws, data_start_row=3, max_col=preg_max_col)
        if preg_removed_count > 0:
            print(f"INFO: Removed {preg_removed_count} formula-only rows from {pregnancy_source_sheet_name}")

        preg_code_col_idx = _find_pregnance_code_column(preg_headers)
        if preg_code_col_idx is None:
            print("WARNING: Could not find Pregnance_code column in EPI-Pregnancy sheet.")
        else:
            _warn_code_quality(preg_ws, preg_code_col_idx, "pw_code", preg_rows_to_keep)

        preg_out_ws = out_wb.create_sheet(title=pregnancy_target_sheet_name)
        preg_out_ws.append(preg_headers)
        for row_num in preg_rows_to_keep:
            row_values = [preg_ws.cell(row=row_num, column=col_idx).value for col_idx in range(1, preg_max_col + 1)]
            preg_out_ws.append(row_values)
    else:
        print(f"WARNING: Pregnancy source sheet not found: {pregnancy_source_sheet_name}")

    try:
        out_wb.save(output_path)
    except Exception as exc:
        raise RuntimeError(f"Failed to save workbook: {exc}") from exc

    print(f"Done: created {output_path.name}")
    return 0


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Create CHDN_EPI_clean.xlsx with completion columns")
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    sys.exit(build_clean_sheet(args.input, args.output))
